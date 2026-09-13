from __future__ import annotations

import logging

from app.domain.input_limits import validate_pdf_page_count

logger = logging.getLogger(__name__)
MAX_EXTRACTED_TEXT_CHARS = 100_000


class TextExtractionService:
    """Service for extracting text from various file formats."""

    @staticmethod
    def _bounded_text(text: str) -> str:
        """Bound persisted extraction text to protect DB and worker memory."""
        return text[:MAX_EXTRACTED_TEXT_CHARS]

    @staticmethod
    def _append_bounded_part(parts: list[str], value: str, extracted_chars: int) -> int:
        """Append one extracted fragment without building an oversized buffer."""
        separator = "\n" if parts else ""
        remaining = MAX_EXTRACTED_TEXT_CHARS - extracted_chars - len(separator)
        if remaining <= 0:
            return extracted_chars
        bounded = value[:remaining]
        if bounded:
            parts.append(separator + bounded)
            return extracted_chars + len(separator) + len(bounded)
        return extracted_chars

    async def extract_text(self, file_path: str, file_type: str) -> str:
        """Extract text from a file based on its type.

        Args:
            file_path: Path to the file.
            file_type: File extension (pdf, docx, jpg, etc.).

        Returns:
            Extracted text content.
        """
        if file_type in ("pdf",):
            return await self._extract_pdf(file_path)
        elif file_type == "docx":
            return await self._extract_docx(file_path)
        elif file_type in ("jpg", "jpeg", "png"):
            return await self._extract_image(file_path)
        elif file_type == "eml":
            return await self._extract_email(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    async def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF using pdfplumber."""
        try:
            import pdfplumber

            text_parts = []
            with pdfplumber.open(file_path) as pdf:
                pages = pdf.pages
                validate_pdf_page_count(len(pages))
                extracted_chars = 0
                for page in pages:
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
                    page_text = page.extract_text()
                    if page_text:
                        remaining = MAX_EXTRACTED_TEXT_CHARS - extracted_chars
                        bounded_page_text = page_text[:remaining]
                        text_parts.append(bounded_page_text)
                        extracted_chars += len(bounded_page_text)
            return self._bounded_text("\n\n".join(text_parts))
        except Exception as e:
            logger.error("PDF extraction failed (%s)", type(e).__name__)
            raise

    async def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts: list[str] = []
            extracted_chars = 0
            for paragraph in doc.paragraphs:
                paragraph_text = paragraph.text
                if paragraph_text.strip():
                    extracted_chars = self._append_bounded_part(
                        text_parts, paragraph_text, extracted_chars
                    )
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
            for table in doc.tables:
                if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                    break
                for row in table.rows:
                    row_started = False
                    for cell in row.cells:
                        cell_text = cell.text
                        if not cell_text.strip():
                            continue
                        separator = " | " if row_started else "\n"
                        if not text_parts:
                            separator = ""
                        remaining = MAX_EXTRACTED_TEXT_CHARS - extracted_chars - len(separator)
                        if remaining <= 0:
                            break
                        bounded = cell_text[:remaining]
                        if bounded:
                            text_parts.append(separator + bounded)
                            extracted_chars += len(separator) + len(bounded)
                        row_started = True
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
            return "".join(text_parts)
        except Exception as e:
            logger.error("DOCX extraction failed (%s)", type(e).__name__)
            raise

    async def _extract_image(self, file_path: str) -> str:
        """Extract text from image - returns path for vision model processing."""
        return f"[IMAGE:{file_path}]"

    async def _extract_email(self, file_path: str) -> str:
        """Extract text from email file."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(MAX_EXTRACTED_TEXT_CHARS)
        except Exception as e:
            logger.error("Email extraction failed (%s)", type(e).__name__)
            raise

    def is_image(self, file_type: str) -> bool:
        """Check if file type is an image."""
        return file_type in ("jpg", "jpeg", "png")

    def extract_text_sync(self, file_path: str, file_type: str) -> str:
        """Extract text from a file synchronously (for Celery tasks)."""
        if file_type in ("pdf",):
            return self._extract_pdf_sync(file_path)
        elif file_type == "docx":
            return self._extract_docx_sync(file_path)
        elif file_type in ("jpg", "jpeg", "png"):
            return f"[IMAGE:{file_path}]"
        elif file_type == "eml":
            return self._extract_email_sync(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    def _extract_pdf_sync(self, file_path: str) -> str:
        """Extract text from PDF synchronously."""
        try:
            import pdfplumber

            text_parts = []
            with pdfplumber.open(file_path) as pdf:
                pages = pdf.pages
                validate_pdf_page_count(len(pages))
                extracted_chars = 0
                for page in pages:
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
                    page_text = page.extract_text()
                    if page_text:
                        remaining = MAX_EXTRACTED_TEXT_CHARS - extracted_chars
                        bounded_page_text = page_text[:remaining]
                        text_parts.append(bounded_page_text)
                        extracted_chars += len(bounded_page_text)
            return self._bounded_text("\n\n".join(text_parts))
        except Exception as e:
            logger.error("PDF extraction failed (%s)", type(e).__name__)
            raise

    def _extract_docx_sync(self, file_path: str) -> str:
        """Extract text from DOCX synchronously."""
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts: list[str] = []
            extracted_chars = 0
            for paragraph in doc.paragraphs:
                paragraph_text = paragraph.text
                if paragraph_text.strip():
                    extracted_chars = self._append_bounded_part(
                        text_parts, paragraph_text, extracted_chars
                    )
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
            for table in doc.tables:
                if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                    break
                for row in table.rows:
                    row_started = False
                    for cell in row.cells:
                        cell_text = cell.text
                        if not cell_text.strip():
                            continue
                        separator = " | " if row_started else "\n"
                        if not text_parts:
                            separator = ""
                        remaining = MAX_EXTRACTED_TEXT_CHARS - extracted_chars - len(separator)
                        if remaining <= 0:
                            break
                        bounded = cell_text[:remaining]
                        if bounded:
                            text_parts.append(separator + bounded)
                            extracted_chars += len(separator) + len(bounded)
                        row_started = True
                    if extracted_chars >= MAX_EXTRACTED_TEXT_CHARS:
                        break
            return "".join(text_parts)
        except Exception as e:
            logger.error("DOCX extraction failed (%s)", type(e).__name__)
            raise

    def _extract_email_sync(self, file_path: str) -> str:
        """Extract text from email file synchronously."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(MAX_EXTRACTED_TEXT_CHARS)
        except Exception as e:
            logger.error("Email extraction failed (%s)", type(e).__name__)
            raise


text_extraction_service = TextExtractionService()
